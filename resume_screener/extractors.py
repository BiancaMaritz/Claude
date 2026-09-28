"""Text extraction for PDF, DOCX and TXT resumes."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")

# A PDF yielding less text than this is treated as scanned/image-only and is
# sent to Claude as a native PDF document block instead of extracted text.
MIN_PDF_TEXT_CHARS = 200


class ExtractionError(Exception):
    """Raised when a resume file cannot be read."""


@dataclass
class ResumeDocument:
    path: Path
    text: str
    # Raw bytes are kept only for PDFs whose text layer is missing or thin.
    pdf_bytes: bytes | None = None

    @property
    def filename(self) -> str:
        return self.path.name

    @property
    def use_native_pdf(self) -> bool:
        return self.pdf_bytes is not None


def _read_txt(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError(f"Unable to decode text file: {path.name}")


def _read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("python-docx is not installed") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Invalid or corrupt DOCX file: {path.name} ({exc})") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Many resume templates put skills and history inside tables.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _read_pdf(path: Path) -> tuple[str, bytes]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("pypdf is not installed") from exc

    data = path.read_bytes()
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError(f"PDF is password protected: {path.name}") from exc
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Invalid or corrupt PDF file: {path.name} ({exc})") from exc
    return text, data


def extract_resume(path: str | Path) -> ResumeDocument:
    path = Path(path)
    if not path.is_file():
        raise ExtractionError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file type '{suffix}' for {path.name}")

    if suffix == ".pdf":
        text, data = _read_pdf(path)
        text = text.strip()
        if len(text) < MIN_PDF_TEXT_CHARS:
            return ResumeDocument(path=path, text=text, pdf_bytes=data)
        return ResumeDocument(path=path, text=text)

    text = (_read_docx(path) if suffix == ".docx" else _read_txt(path)).strip()
    if not text:
        raise ExtractionError(f"No readable text found in {path.name}")
    return ResumeDocument(path=path, text=text)
