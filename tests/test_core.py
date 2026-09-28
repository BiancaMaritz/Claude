import csv
from types import SimpleNamespace

import anthropic
import httpx2 as httpx
import pytest

from resume_screener.analyzer import FALLBACK_BETA, CandidateReport, ResumeAnalyzer, ScreeningResult
from resume_screener.extractors import ExtractionError, extract_resume
from resume_screener.reports import export_csv, export_pdf, format_text_report, ranked


def make_result(match=90, name="Jane Doe"):
    return ScreeningResult(
        candidate_name=name,
        match_percentage=match,
        matched_skills=["Python", "SQL"],
        missing_skills=["Kubernetes"],
        relevant_experience_years=6,
        experience_fit="Strong fit.",
        strengths=["Delivery"],
        gaps=["No cloud certs"],
        interview_questions=["Describe a migration you led."],
        summary="Recommended.",
    )


class FakeMessages:
    def __init__(self, response=None, exc=None):
        self.response, self.exc, self.calls = response, exc, []

    def parse(self, **kwargs):
        self.calls.append(kwargs)
        if self.exc:
            raise self.exc
        return self.response


def fake_client(**kw):
    messages = FakeMessages(**kw)
    return SimpleNamespace(beta=SimpleNamespace(messages=messages)), messages


# ---------- extractors ----------
def test_extract_txt(tmp_path):
    f = tmp_path / "cv.txt"
    f.write_text("Jane Doe\nPython developer", encoding="utf-8")
    assert "Python developer" in extract_resume(f).text


def test_extract_docx_with_table(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("John Smith")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Skills", "Automation Anywhere"
    d.save(tmp_path / "cv.docx")
    text = extract_resume(tmp_path / "cv.docx").text
    assert "John Smith" in text and "Automation Anywhere" in text


def test_extract_pdf_text_and_scanned_fallback(tmp_path):
    from reportlab.pdfgen import canvas

    rich = tmp_path / "rich.pdf"
    c = canvas.Canvas(str(rich))
    for i in range(20):
        c.drawString(50, 800 - i * 20, f"Line {i}: experienced data analyst with SQL and Python skills")
    c.save()
    doc = extract_resume(rich)
    assert not doc.use_native_pdf and "data analyst" in doc.text

    thin = tmp_path / "thin.pdf"
    c = canvas.Canvas(str(thin))
    c.drawString(50, 800, "x")
    c.save()
    assert extract_resume(thin).use_native_pdf


def test_extract_errors(tmp_path):
    with pytest.raises(ExtractionError):
        extract_resume(tmp_path / "missing.pdf")
    bad = tmp_path / "cv.rtf"
    bad.write_text("x")
    with pytest.raises(ExtractionError):
        extract_resume(bad)
    corrupt = tmp_path / "bad.docx"
    corrupt.write_bytes(b"not a zip")
    with pytest.raises(ExtractionError):
        extract_resume(corrupt)


# ---------- analyzer ----------
def _doc(tmp_path):
    f = tmp_path / "cv.txt"
    f.write_text("Jane Doe, Python, SQL, 6 years")
    return extract_resume(f)


def test_analyze_success_and_request_shape(tmp_path):
    resp = SimpleNamespace(stop_reason="end_turn", parsed_output=make_result(85), model="claude-opus-5")
    client, msgs = fake_client(response=resp)
    report = ResumeAnalyzer("k", client=client).analyze(_doc(tmp_path), "Need Python and SQL")
    assert report.status == "Selected" and report.match == 85
    call = msgs.calls[0]
    assert call["model"] == "claude-opus-5"
    assert call["output_format"] is ScreeningResult
    assert call["fallbacks"] == "default" and call["betas"] == [FALLBACK_BETA]
    assert "Need Python and SQL" in call["system"][1]["text"]
    assert call["system"][1]["cache_control"] == {"type": "ephemeral"}


def test_threshold_boundary():
    assert CandidateReport("a", result=make_result(84)).status == "Rejected"
    assert CandidateReport("a", result=make_result(85)).status == "Selected"
    assert CandidateReport("a", result=make_result(140)).match == 100
    assert CandidateReport("a", error="x").status == "Error"


def test_sonnet_has_no_fallbacks(tmp_path):
    resp = SimpleNamespace(stop_reason="end_turn", parsed_output=make_result(), model="claude-sonnet-5")
    client, msgs = fake_client(response=resp)
    ResumeAnalyzer("k", model="claude-sonnet-5", client=client).analyze(_doc(tmp_path), "JD")
    assert "fallbacks" not in msgs.calls[0]


def test_refusal_and_auth_error(tmp_path):
    client, _ = fake_client(response=SimpleNamespace(stop_reason="refusal", parsed_output=None, model="m"))
    assert "declined" in ResumeAnalyzer("k", client=client).analyze(_doc(tmp_path), "JD").error

    req = httpx.Request("POST", "https://api.anthropic.com/v1/messages")
    exc = anthropic.AuthenticationError("bad", response=httpx.Response(401, request=req), body=None)
    client, _ = fake_client(exc=exc)
    assert ResumeAnalyzer("k", client=client).analyze(_doc(tmp_path), "JD").error.startswith("Invalid API key")


# ---------- reports ----------
def _reports():
    return [
        CandidateReport("low.pdf", result=make_result(60, "Low")),
        CandidateReport("err.pdf", error="boom"),
        CandidateReport("high.pdf", result=make_result(92, "High")),
    ]


def test_ranked_and_text():
    order = [r.filename for r in ranked(_reports())]
    assert order == ["high.pdf", "low.pdf", "err.pdf"]
    text = format_text_report(_reports()[2])
    assert "Match: 92%" in text and "Selected" in text and "Kubernetes" in text


def test_export_csv_and_pdf(tmp_path):
    export_csv(_reports(), tmp_path / "r.csv")
    with open(tmp_path / "r.csv", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert rows[0]["Candidate"] == "High" and rows[0]["Status"] == "Selected"
    assert rows[2]["Error"] == "boom"

    export_pdf(_reports(), tmp_path / "r.pdf", job_title="Analyst <Senior> & Co")
    assert (tmp_path / "r.pdf").read_bytes().startswith(b"%PDF")
