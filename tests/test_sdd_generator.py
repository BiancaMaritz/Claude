import json
import os
import zipfile

import pytest

from sdd_generator.__main__ import main
from sdd_generator.analyze import analyse, sql_facts
from sdd_generator.document import build
from sdd_generator.loader import load_export
from sdd_generator.render import to_markdown

R = "Automation Anywhere/Bots/"


def var(name, typ="STRING", inp=False, out=False, default=""):
    return {"name": name, "type": typ, "input": inp, "output": out,
            "defaultValue": {"type": typ, "string": default}}


def attrs(**kw):
    return [{"name": k, "value": {"type": "STRING", "expression": v}} for k, v in kw.items()]


def run_task(path):
    return {"commandName": "runTask", "packageName": "TaskBot", "attributes": [
        {"name": "taskbot", "value": {"type": "TASKBOT", "taskbotFile": {
            "type": "FILE", "string": "repository:///" + path.replace(" ", "%20")}}}]}


def step(label, children=()):
    return {"commandName": "step", "packageName": "Step", "children": list(children),
            "attributes": attrs(name=label)}


def sql(q):
    return {"commandName": "insertUpdateDelete", "packageName": "Database", "attributes": attrs(query=q)}


MAIN = R + "Acme/Logic Manager/Main"
WORK = R + "Acme/Work/Worker"
LOG = R + "Task Logging/Log"
LOOP_A = R + "Acme/Work/LoopA"
LOOP_B = R + "Acme/Work/LoopB"


def bot(nodes, variables=()):
    return {"nodes": nodes, "variables": list(variables), "triggers": [], "packages": []}


BOTS = {
    MAIN: bot([step("Connect", [sql("UPDATE $Db$.[dbo].[Jobs] SET [Status] = 'X', [Tries] = 1 WHERE [Id] = 2")]),
               step("Do work", [run_task(WORK)]), run_task(LOG)],
              [var("Environment", inp=True, default="Prod"), var("DbPassword", inp=True, default="hunter2")]),
    WORK: bot([step("Fetch", [{"commandName": "restPost", "packageName": "Rest",
                              "attributes": attrs(uri="$URL$/api/x")}]),
               {"commandName": "messageBox", "packageName": "MessageBox", "attributes": []}]),
    LOG: bot([sql("INSERT INTO [dbo].[RunLog] ([Name],[At]) VALUES ('a', 1)")]),
    LOOP_A: bot([run_task(LOOP_B)]),
    LOOP_B: bot([run_task(LOOP_A)]),
}


@pytest.fixture
def export_dir(tmp_path):
    for p, d in BOTS.items():
        f = tmp_path / p
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(json.dumps(d))
    (tmp_path / "manifest.json").write_text(json.dumps({
        "files": [{"path": p.replace("/", "\\"), "contentType": "application/vnd.aa.taskbot",
                   "scannedDependencies": [], "description": ""} for p in BOTS],
        "packages": [{"name": "String", "version": "1"}, {"name": "String", "version": "2"}],
        "globalValues": []}))
    return tmp_path


def test_sql_facts():
    ops, cols = sql_facts("UPDATE $Db$.[dbo].[Jobs] SET [Status]='a', [Tries]=1 WHERE Id=2")
    assert ops == [("Jobs", "update")] and cols["Jobs"] == {"Status", "Tries"}
    ops, cols = sql_facts("INSERT INTO [dbo].[RunLog] ([Name],[At]) VALUES ('a',1)")
    assert cols["RunLog"] == {"Name", "At"}
    assert sql_facts("send a note from the server")[0] == []


def test_analysis(export_dir):
    ex = load_export(str(export_dir))
    a = analyse(ex)
    assert len(ex.bots) == 5
    assert a["ref"][MAIN] == "Main Bot 1"
    assert a["ref"][LOG] == "Common Bot 1"
    assert MAIN in a["entry"] and WORK not in a["entry"]
    assert [sorted(c) for c in a["cycles"]] == [sorted([LOOP_A, LOOP_B])]
    assert set(a["entry"]) >= {LOOP_A, LOOP_B}  # a cycle nobody else calls is still an entry group
    assert set(a["tables"]) == {"Jobs", "RunLog"}
    assert a["rest"][0]["method"] == "POST" and a["rest"][0]["uri"] == "{URL}/api/x"
    cats = {h[0] for h in a["hygiene"]}
    assert {"Circular call", "Unattended risk", "Package versions"} <= cats


def test_zip_matches_directory(export_dir, tmp_path):
    z = tmp_path.parent / "x.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for root, _, files in os.walk(export_dir):
            for f in files:
                full = os.path.join(root, f)
                zf.write(full, os.path.relpath(full, export_dir))
    assert set(load_export(str(z)).bots) == set(load_export(str(export_dir)).bots)


def test_secret_defaults_masked_and_intake(export_dir):
    ex = load_export(str(export_dir))
    md = to_markdown(build(ex, analyse(ex), {"process_description": "Does things.",
                                            "assumptions": ["A1"]}, title="T", date="2026-01-01"))
    assert "hunter2" not in md and "[masked]" in md
    assert "Does things." in md and "- A1" in md
    assert "[TO BE COMPLETED]: functional requirements." in md


def test_cli_outputs(export_dir, tmp_path):
    out = tmp_path / "out"
    assert main([str(export_dir), "--out", str(out), "--date", "2026-01-01"]) == 0
    assert (out / "SDD.md").read_text().startswith("# ")
    from docx import Document
    assert len(Document(str(out / "SDD.docx")).tables) > 5
    assert json.loads((out / "SDD.facts.json").read_text())["entry_points"]


def test_cli_with_intake_stakeholders(export_dir, tmp_path):
    intake = tmp_path / "intake.json"
    intake.write_text(json.dumps({"stakeholders": [
        {"team": "Team A", "name": "N", "role": "R", "email": "e", "approval": "High"}]}))
    out = tmp_path / "o"
    assert main([str(export_dir), "--out", str(out), "--intake", str(intake)]) == 0
    from docx import Document
    for t in Document(str(out / "SDD.docx")).tables:
        assert len({len(r.cells) for r in t.rows}) == 1


def zip_of(export_dir, skip=()):
    import io
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zf:
        for root, _, files in os.walk(export_dir):
            for f in files:
                full = os.path.join(root, f)
                rel = os.path.relpath(full, export_dir)
                if rel.replace(os.sep, "/") not in skip:
                    zf.write(full, rel)
    return z.getvalue()


def test_store_add_vs_overwrite(export_dir, tmp_path):
    from sdd_generator.store import Store, deep_merge
    from sdd_generator.web import classify_uploads

    store = Store(str(tmp_path / "docs"))
    full = classify_uploads([("full.zip", zip_of(export_dir))])[0]
    part = classify_uploads([("part.zip", zip_of(export_dir, skip=("Automation Anywhere/Bots/Task Logging/Log", "manifest.json")))])[0]
    r = store.load("d", "add", part, [], ["part.zip"])
    assert (r["added"], r["created"]) == (4, True)
    r = store.load("d", "add", full, [], ["full.zip"])
    assert (r["added"], r["unchanged"], r["removed"]) == (1, 4, 0)
    assert len(store.get("d")["export"]["bots"]) == 5
    r = store.load("d", "overwrite", part, [], ["part.zip"])
    assert (r["removed"], r["added"], r["unchanged"]) == (1, 0, 4)
    assert len(store.get("d")["export"]["bots"]) == 4
    assert deep_merge({"a": [1], "b": 1}, {"a": [1, 2], "b": 2}) == {"a": [1, 2], "b": 2}


def test_add_keeps_logic_when_manifest_only_arrives(export_dir, tmp_path):
    from sdd_generator.store import Store, export_from_dict
    from sdd_generator.web import classify_uploads

    store = Store(str(tmp_path / "docs"))
    store.load("d", "add", classify_uploads([("a.zip", zip_of(export_dir))])[0], [], ["a.zip"])
    only_manifest = zip_of(export_dir, skip=tuple(BOTS))
    store.load("d", "add", classify_uploads([("m.zip", only_manifest)])[0], [], ["m.zip"])
    ex = export_from_dict(store.get("d")["export"])
    assert all(b.data is not None for b in ex.bots.values())


def test_web_api(export_dir, tmp_path):
    import threading
    import urllib.error
    import urllib.request
    from http.server import ThreadingHTTPServer

    from sdd_generator.web import App, make_handler

    srv = ThreadingHTTPServer(("127.0.0.1", 0), None)
    port = srv.server_address[1]
    srv.RequestHandlerClass = make_handler(App(str(tmp_path / "docs")), port)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{port}"
    b = "XBX"

    def post(fields, files):
        body = b""
        for k, v in fields.items():
            body += f'--{b}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode()
        for fname, data in files:
            body += (f'--{b}\r\nContent-Disposition: form-data; name="files"; filename="{fname}"\r\n'
                     f'Content-Type: application/octet-stream\r\n\r\n').encode() + data + b"\r\n"
        body += f"--{b}--\r\n".encode()
        req = urllib.request.Request(base + "/api/load", body,
                                     {"Content-Type": f"multipart/form-data; boundary={b}"})
        return json.loads(urllib.request.urlopen(req).read())

    try:
        assert b"SDD Generator" in urllib.request.urlopen(base + "/").read()
        j = post({"doc": "My Doc", "mode": "add", "title": "T"}, [("a b.zip", zip_of(export_dir)),
                                                                  ("intake.json", b'{"assumptions": ["A1"]}')])
        assert j["report"]["added"] == 5 and j["report"]["intake"] == 1
        d = json.loads(urllib.request.urlopen(base + "/api/doc?name=My%20Doc").read())
        assert d["title"] == "T" and "A1" in d["html"] and d["menu"][0]["text"].startswith("1.")
        assert urllib.request.urlopen(base + "/download/My%20Doc/docx").read()[:2] == b"PK"
        assert b"# T" in urllib.request.urlopen(base + "/download/My%20Doc/md").read()
        for bad in ([("x.txt", b"hi")], [("bad.zip", b"not a zip")], [("m.json", b'{"files": []}')]):
            with pytest.raises(urllib.error.HTTPError) as e:
                post({"doc": "My Doc", "mode": "add"}, bad)
            assert e.value.code == 400
        with pytest.raises(urllib.error.HTTPError):  # path traversal in the name
            post({"doc": "../evil", "mode": "add"}, [("a.zip", zip_of(export_dir))])
        req = urllib.request.Request(base + "/api/delete", b'{"name": "My Doc"}', {"Origin": "http://evil.example"})
        with pytest.raises(urllib.error.HTTPError) as e:  # cross-site request is refused
            urllib.request.urlopen(req)
        assert e.value.code == 403
        assert json.loads(urllib.request.urlopen(base + "/api/docs").read())[0]["bots"] == 5
    finally:
        srv.shutdown()


def test_reload_same_content_is_detected(export_dir, tmp_path):
    from sdd_generator.store import Store
    from sdd_generator.web import classify

    store = Store(str(tmp_path / "docs"))
    z = zip_of(export_dir)
    ex, it, names, src = classify([("v1.zip", z)], [])
    r1 = store.load("d", "add", ex, it, names, sources=src)
    assert not r1["noop"] and r1["already_loaded"] == []
    n_hist = len(store.get("d")["history"])
    # same content again, even under another file name: detected, nothing written
    ex, it, names, src = classify([("renamed copy.zip", z)], [])
    r2 = store.load("d", "add", ex, it, names, sources=src)
    assert r2["noop"] and r2["unchanged"] == 5 and r2["already_loaded"] == ["renamed copy.zip"]
    assert len(store.get("d")["history"]) == n_hist
    # overwrite with identical content is also a no-op
    assert store.load("d", "overwrite", ex, it, names, sources=src)["noop"]


def test_reload_older_version_still_applies(export_dir, tmp_path):
    from sdd_generator.store import Store
    from sdd_generator.web import classify

    store = Store(str(tmp_path / "docs"))
    v1 = zip_of(export_dir)
    changed = json.loads((export_dir / WORK).read_text())
    changed["nodes"].append(step("New step"))
    (export_dir / WORK).write_text(json.dumps(changed))
    v2 = zip_of(export_dir)
    for name, z in (("v1.zip", v1), ("v2.zip", v2), ("v1.zip", v1)):
        ex, it, names, src = classify([(name, z)], [])
        r = store.load("d", "add", ex, it, names, sources=src)
    # v1 was loaded before, but the document now holds v2's Worker, so reloading v1 must revert it
    assert r["already_loaded"] == ["v1.zip"] and not r["noop"] and r["updated"] == 1
    assert [n.rsplit("/", 1)[-1] for n in r["names"]["updated"]] == ["Worker"]


def test_zip_fingerprint_ignores_name_and_location(export_dir, tmp_path):
    a, b = tmp_path / "a.zip", tmp_path / "sub" / "b.zip"
    b.parent.mkdir()
    data = zip_of(export_dir)
    a.write_bytes(data)
    b.write_bytes(data)
    assert load_export(str(a)).fingerprint == load_export(str(b)).fingerprint != ""
    (export_dir / "Automation Anywhere/Bots/lib.jar").write_bytes(b"x")
    b.write_bytes(zip_of(export_dir, skip=("a.zip", "sub/b.zip")))
    assert load_export(str(a)).fingerprint != load_export(str(b)).fingerprint  # an added file changes it


def test_binaries_are_never_read(export_dir, tmp_path, monkeypatch):
    (export_dir / "Automation Anywhere/Bots/huge.jar").write_bytes(b"\0" * 5_000_000)
    z = tmp_path / "x.zip"
    z.write_bytes(zip_of(export_dir))
    read = []
    orig = zipfile.ZipFile.read
    monkeypatch.setattr(zipfile.ZipFile, "read", lambda self, n, *a, **k: (read.append(getattr(n, "filename", n)), orig(self, n, *a, **k))[1])
    assert len(load_export(str(z)).bots) == 5
    assert not any(str(n).endswith(".jar") for n in read)


def test_web_load_from_path(export_dir, tmp_path):
    import threading
    import urllib.request
    from http.server import ThreadingHTTPServer

    from sdd_generator.web import App, make_handler

    z = tmp_path / "big export.zip"
    z.write_bytes(zip_of(export_dir))
    srv = ThreadingHTTPServer(("127.0.0.1", 0), None)
    port = srv.server_address[1]
    srv.RequestHandlerClass = make_handler(App(str(tmp_path / "docs")), port)
    threading.Thread(target=srv.serve_forever, daemon=True).start()

    def load(path):
        b = "XB"
        body = "".join(f'--{b}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'
                       for k, v in {"doc": "P", "mode": "add", "paths": path}.items()) + f"--{b}--\r\n"
        req = urllib.request.Request(f"http://127.0.0.1:{port}/api/load", body.encode(),
                                     {"Content-Type": f"multipart/form-data; boundary={b}"})
        try:
            return json.loads(urllib.request.urlopen(req).read())
        except urllib.error.HTTPError as e:
            return json.loads(e.read())

    try:
        assert load(f'"{z}"')["report"]["added"] == 5          # Windows "Copy as path" adds quotes
        assert load(str(z))["report"]["noop"] is True          # same file again
        assert load(str(export_dir))["report"]["already_loaded"] == []  # a folder is a different source
        assert "not found" in load(str(tmp_path / "nope.zip"))["error"]
        (tmp_path / "notes.txt").write_text("hi")
        assert "give a .zip" in load(str(tmp_path / "notes.txt"))["error"]
    finally:
        srv.shutdown()
