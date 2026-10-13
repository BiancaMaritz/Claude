"""Render document blocks to Markdown and Word (.docx)."""
from __future__ import annotations

import os
import re

BOLD = re.compile(r"\*\*(.+?)\*\*")
PLACEHOLDER = re.compile(r"(\[(?:TO BE COMPLETED|CONFIRM|Screenshot)[^\]]*\])")


def _lines(cell):
    if isinstance(cell, (list, tuple)):
        return [str(x) for x in cell] or [""]
    return str(cell).split("\n")


# --------------------------------------------------------------- markdown ---

def _md_cell(cell):
    return "<br>".join(_lines(cell)).replace("|", "\\|") or " "


def to_markdown(blocks) -> str:
    out = []
    headings = [(b[1], b[2]) for b in blocks if b[0] == "h" and b[1] <= 2]
    for b in blocks:
        k = b[0]
        if k == "title":
            out.append(f"# {b[1]}\n")
        elif k == "toc":
            out.append("## Contents\n")
            for lvl, t in headings:
                out.append(f"{'  ' * (lvl - 1)}- [{t}](#{_slug(t)})")
            out.append("")
        elif k == "h":
            out.append(f"{'#' * (b[1] + 1)} {b[2]}\n")
        elif k in ("p", "todo"):
            out.append(f"{b[1]}\n")
        elif k == "src":
            out.append(f"*{b[1]}*\n")
        elif k == "img":
            out.append(f"![{b[2]}](media/{b[1]['file']})\n\n*{b[2]}*\n")
        elif k == "ul":
            out += [f"- {t}" for t in b[1]] + [""]
        elif k == "table":
            _, head, rows = b
            out.append("| " + " | ".join(head) + " |")
            out.append("|" + "|".join(" --- " for _ in head) + "|")
            for r in rows:
                out.append("| " + " | ".join(_md_cell(c) for c in r) + " |")
            out.append("")
    return "\n".join(out)


def _slug(t):
    return re.sub(r"[^\w\- ]", "", t.lower()).replace(" ", "-")


# ------------------------------------------------------------------- docx ---

def _add_runs(par, text, bold=False):
    from docx.enum.text import WD_COLOR_INDEX

    for part in PLACEHOLDER.split(text):
        if not part:
            continue
        if PLACEHOLDER.fullmatch(part):
            r = par.add_run(part)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = bold
            continue
        for i, seg in enumerate(BOLD.split(part)):
            if seg:
                r = par.add_run(seg)
                r.bold = bold or i % 2 == 1


def _toc_field(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    par = doc.add_paragraph()
    run = par.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                el.set(qn("w:dirty"), "true")
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run2 = par.add_run("Right-click and choose Update Field to build the table of contents.")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run2._r.append(end)


def _fit_columns(table, head, rows, total=6.9):
    """Size columns by typical content length so short ID columns stay narrow."""
    from docx.shared import Inches

    def weight(i):
        lens = [max(len(x) for x in _lines(r[i])) for r in rows if i < len(r)] or [6]
        lens.append(len(head[i]))
        return min(max(sum(lens) / len(lens), 8), 60)

    w = [weight(i) for i in range(len(head))]
    table.autofit = False
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(total * w[i] / sum(w))


def to_docx(blocks, path):
    from docx import Document
    from docx.shared import Pt

    from docx.shared import Inches, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for b in blocks:
        k = b[0]
        if k == "title":
            doc.add_heading(b[1], 0)
        elif k == "toc":
            doc.add_heading("Contents", 1)
            _toc_field(doc)
        elif k == "h":
            doc.add_heading(b[2], b[1])
        elif k in ("p", "todo"):
            _add_runs(doc.add_paragraph(), b[1])
        elif k == "src":
            r = doc.add_paragraph().add_run(b[1])
            r.italic = True
            r.font.color.rgb = RGBColor(0x1F, 0x5F, 0xBF)
        elif k == "img":
            if os.path.exists(b[1]["path"]):
                doc.add_picture(b[1]["path"], width=Inches(6.2))
            else:
                _add_runs(doc.add_paragraph(), f"[Image file missing: {b[1]['file']}]")
            cap = doc.add_paragraph().add_run(b[2])
            cap.italic = True
        elif k == "ul":
            for t in b[1]:
                _add_runs(doc.add_paragraph(style="List Bullet"), t)
        elif k == "table":
            _, head, rows = b
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            for i, h in enumerate(head):
                c = t.rows[0].cells[i]
                c.text = ""
                _add_runs(c.paragraphs[0], h, bold=True)
            for r in rows:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    lines = _lines(v)
                    cells[i].text = ""
                    _add_runs(cells[i].paragraphs[0], lines[0])
                    for ln in lines[1:]:
                        _add_runs(cells[i].add_paragraph(), ln)
            _fit_columns(t, head, rows)
            doc.add_paragraph()
    doc.save(path)


# ------------------------------------------------------------------- html ---

def _h(text) -> str:
    """Escape, then apply **bold** and highlight placeholders."""
    import html as _html

    t = _html.escape(str(text))
    t = BOLD.sub(r"<strong>\1</strong>", t)
    return re.sub(r"(\[(?:TO BE COMPLETED|CONFIRM|Screenshot)[^\]]*\])", r'<span class="todo">\1</span>', t)


def to_html(blocks):
    """Return (menu, html). menu: [{id, level, text, todo}] for h1/h2 headings.

    todo counts the placeholders still open in that section, so the menu shows what is left.
    """
    import html as _html

    menu, out = [], []
    cur = {1: None, 2: None}
    n = 0

    def cell(c):
        return "<br>".join(_h(x) for x in _lines(c))

    for b in blocks:
        k = b[0]
        todo = 0
        if k == "title":
            out.append(f'<h1 class="doc-title">{_h(b[1])}</h1>')
        elif k == "h":
            n += 1
            lvl, text = b[1], b[2]
            hid = f"s{n}"
            out.append(f'<h{lvl + 1} id="{hid}">{_h(text)}</h{lvl + 1}>')
            if lvl <= 2:
                entry = {"id": hid, "level": lvl, "text": text, "todo": 0, "docs": 0}
                menu.append(entry)
                cur[lvl] = entry
                if lvl == 1:
                    cur[2] = None
        elif k == "src":
            out.append(f'<div class="src">{_h(b[1])}</div>')
            for lvl in (1, 2):
                if cur[lvl]:
                    cur[lvl]["docs"] = cur[lvl].get("docs", 0) + 1
        elif k == "img":
            shown = b[1]["svg"] if b[1].get("svg") else f'<img src="{_html.escape(b[1]["url"], quote=True)}" alt="{_html.escape(b[2], quote=True)}">'
            out.append(f'<figure class="{"wide" if b[1].get("svg") else ""}">{shown}<figcaption>{_h(b[2])}</figcaption></figure>')
        elif k in ("p", "todo"):
            out.append(f"<p>{_h(b[1])}</p>")
            todo = len(PLACEHOLDER.findall(str(b[1])))
        elif k == "ul":
            out.append("<ul>" + "".join(f"<li>{_h(t)}</li>" for t in b[1]) + "</ul>")
            todo = sum(len(PLACEHOLDER.findall(str(t))) for t in b[1])
        elif k == "table":
            _, head, rows = b
            body = "".join("<tr>" + "".join(f"<td>{cell(c)}</td>" for c in r) + "</tr>" for r in rows)
            out.append('<div class="tw"><table><thead><tr>' + "".join(f"<th>{_h(x)}</th>" for x in head)
                       + f"</tr></thead><tbody>{body}</tbody></table></div>")
            todo = sum(len(PLACEHOLDER.findall(x)) for r in rows for c in r for x in _lines(c)
                       if "Screenshot" not in x)
        if todo:
            for lvl in (1, 2):
                if cur[lvl]:
                    cur[lvl]["todo"] += todo
    return menu, "\n".join(out)
